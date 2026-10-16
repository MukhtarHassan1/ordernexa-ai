from io import BytesIO

import pytest
from openpyxl import Workbook

from app.services.document_intelligence import extract_document, extract_order_draft
from app.services.source_grounding import inspect_source_text
import app.services.document_intelligence as di


def order_rows():
    return [
        ['PURCHASE ORDER'],
        ['PO Number', 'PO-TABLE-001'],
        ['Customer', 'Northstar Engineering Ltd'],
        ['Ship To', 'Manchester, United Kingdom'],
        ['Requested Delivery Date', '30 October 2026'],
        ['Currency', 'GBP'],
        [],
        ['Line', 'Customer Ref', 'Description', 'Qty', 'UOM', 'Unit Price', 'Line Total'],
        ['1', 'XTR-42', 'Industrial pump assembly', '40', 'EA', '£260.00', '£10,400.00'],
        ['2', 'BRK-90', 'Heavy brake assembly', '15', 'EA', '£165.00', '£2,475.00'],
    ]


def test_xlsx_metadata_and_lines_are_independently_grounded():
    wb = Workbook()
    ws = wb.active
    ws.title = 'PO'
    for row in order_rows():
        ws.append(row)
    buf = BytesIO(); wb.save(buf); wb.close()

    content = extract_document(filename='order.xlsx', data=buf.getvalue())
    assert content.extraction_method == 'XLSX_CELLS'
    source = inspect_source_text(content.text, tables=content.tables)
    assert source.document_type == 'PURCHASE_ORDER'
    assert source.fields['po_number'].value == 'PO-TABLE-001'
    assert source.fields['customer_hint'].value == 'Northstar Engineering Ltd'
    assert source.fields['delivery_warehouse_code'].value == 'MAN'
    assert len(source.lines) == 2
    assert source.lines[0].product_term == 'XTR-42'


@pytest.mark.asyncio
async def test_xlsx_clean_order_uses_fast_path_without_model(monkeypatch):
    wb = Workbook(); ws = wb.active
    for row in order_rows(): ws.append(row)
    buf = BytesIO(); wb.save(buf); wb.close()
    content = extract_document(filename='order.xlsx', data=buf.getvalue())
    monkeypatch.setattr(di, 'get_ai_provider', lambda: (_ for _ in ()).throw(AssertionError('AI should not run')))
    result = await extract_order_draft(content)
    assert result.provider == 'DETERMINISTIC_SOURCE'
    assert result.data.po_number == 'PO-TABLE-001'


def test_csv_metadata_and_lines_are_independently_grounded():
    csv_text = '\n'.join(','.join(row) for row in order_rows())
    content = extract_document(filename='order.csv', data=csv_text.encode())
    assert content.extraction_method == 'CSV_TABLE'
    source = inspect_source_text(content.text, tables=content.tables)
    assert source.fields['po_number'].value == 'PO-TABLE-001'
    assert source.fields['customer_hint'].value == 'Northstar Engineering Ltd'
    assert source.fields['currency'].value == 'GBP'
    assert [line.product_term for line in source.lines] == ['XTR-42', 'BRK-90']
